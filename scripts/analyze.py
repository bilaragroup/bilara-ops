import openpyxl,json,re,datetime,collections,pathlib
ROOT=pathlib.Path(__file__).resolve().parents[1]
out={'purchase':[],'payment':[],'sheets':[]}
def val(v):
 if isinstance(v,(datetime.datetime,datetime.date)):return v.isoformat()[:10]
 return v
def txt(v):return '' if v is None else str(val(v)).strip()
for kind,file in [('purchase','Purchase Plan FY 26-27 (1).xlsx'),('payment','Payment Forecast FY 26-27.xlsx')]:
 w=openpyxl.load_workbook(pathlib.Path('C:/Users/georg/Downloads')/file,data_only=True)
 f=openpyxl.load_workbook(pathlib.Path('C:/Users/georg/Downloads')/file,data_only=False)
 for si,s in enumerate(w):
  headers=[txt(c.value).lower() for c in s[1]]
  count=0; formulas=0;errors=0
  for row in s.iter_rows(min_row=2):
   cells=[val(c.value) for c in row]; raw={openpyxl.utils.get_column_letter(i+1):v for i,v in enumerate(cells) if v is not None}
   if not raw:continue
   for c in row:
    if c.data_type=='e':errors+=1
    if f[s.title][c.coordinate].data_type=='f':formulas+=1
   idx=headers.index('item') if 'item' in headers else (4 if kind=='purchase' else 1)
   if not txt(cells[idx]) or txt(cells[idx]).lower()=='item':continue
   def get(name,default=None):
    ix=next((i for i,h in enumerate(headers) if h==name),None)
    return cells[ix] if ix is not None else default
   r={'id':f'{kind[:3]}-{si+1:02}-{row[0].row:04}','kind':kind,'sheet':s.title,'row':row[0].row,'file':file,'project':txt(cells[0]),'item':txt(cells[idx]),'raw':raw,'links':{c.column_letter:c.hyperlink.target for c in f[s.title][row[0].row] if c.hyperlink and c.hyperlink.target},'vendor':txt(get('vendor',cells[4] if kind=='payment' else None)),'entity':txt(get('wares/builders')),'bill':txt(get('bill'))}
   if kind=='purchase':r.update(boqQty=get('boq quantity'),boqRate=get('boq price'),actualQty=get('purchased quantity'),actualRate=get('actual price'),milestone=txt(get('milestone')),activity=txt(get('activity')),date=get('delivery date',get('requested delivery date')),actualDate=get('actual delivery date'),status=txt(get('purchase status')),approval=txt(get('audit approval')),po=txt(get('po')),grn=txt(get('grn number')),notes=txt(get('remarks')))
   else:r.update(amount=get('payment to be done'),remaining=get('remaining amount (if any)'),date=get('payment date'),status=txt(get('payment status')),receipt=txt(get('payment receipt')))
   r['issues']=[]
   if not re.fullmatch(r'\d+(\.0)?',r['project']):r['issues'].append('Project needs confirmation')
   for k in (['boqQty','boqRate','actualQty','actualRate'] if kind=='purchase' else ['amount']):
    v=r.get(k)
    if v is None or txt(v) in ['','-']:r['issues'].append(k+' missing')
    elif not isinstance(v,(int,float)):r['issues'].append(k+' contains text / units')
    elif abs(v)>10000000:r['issues'].append(k+' unusually large')
   if r.get('date'):r['issues'].append('Confirm date against week (mixed source formats)')
   else:r['issues'].append('Date missing')
   if kind=='payment' and r.get('remaining') not in [None,0,'0']:r['issues'].append('Reconcile remaining balance with payment status')
   out[kind].append(r);count+=1
  out['sheets'].append({'kind':kind,'name':s.title,'rows':count,'formulas':formulas,'errors':errors,'headers':headers})
out['summary']={k:len(out[k]) for k in ['purchase','payment']}
(ROOT/'data').mkdir(exist_ok=True)
(ROOT/'data'/'source.json').write_text(json.dumps(out,ensure_ascii=False,default=str),encoding='utf-8')
print(json.dumps({'summary':out['summary'],'sheets':[{k:v for k,v in x.items() if k!='headers'} for x in out['sheets']]},indent=2))
